#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REFORMATEADOR DE DOCUMENTOS ACADÉMICOS
======================================
Recibe un .docx o .pdf (también .doc/.odt/.rtf si hay LibreOffice instalado) y
genera un .docx nuevo con el formato de la norma elegida: APA 7, IEEE, MLA 9,
Chicago 17, Vancouver o ICONTEC (NTC 1486).

Qué ajusta:  papel y márgenes · fuente · interlineado · sangría · títulos por
nivel · citas en bloque · listas · leyendas de tablas/figuras · lista de
referencias (sangría francesa, orden y numeración opcionales) · numeración de
páginas · encabezado y pie de página.

Principio de diseño: SOLO se toca el formato, nunca el texto.  Para .docx se
edita una copia del documento original (se conservan imágenes, tablas, notas al
pie, campos, hipervínculos).  Al final se compara el texto de entrada con el de
salida y se informa si hay alguna diferencia.

Uso rápido
----------
    python reformateador.py                          # abre la ventana (arrastrar y soltar)
    python reformateador.py tesis.docx -n apa7       # línea de comandos
    python reformateador.py informe.pdf -n ieee -o informe_ieee.docx
    python reformateador.py --listar-normas

Los "presets" de cada norma están en el diccionario NORMAS (más abajo); son una
aproximación razonable de los manuales.  Si su universidad tiene una guía propia
(muy común con ICONTEC), ajuste allí los valores.
"""
from __future__ import annotations

import argparse
import collections
import queue
import re
import shutil
import statistics
import subprocess
import sys
import tempfile
import threading
import unicodedata
from dataclasses import dataclass, field
from pathlib import Path

try:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:  # mensaje claro si falta la dependencia principal
    sys.exit("Falta la librería 'python-docx'.  Instálela con:  pip install -r requirements.txt")


# ════════════════════════════════════════════════════════════════════════════
# 1. DEFINICIÓN DE NORMAS
# ════════════════════════════════════════════════════════════════════════════
def T(alin="izq", b=False, i=False, mayus=None, sang=0.0, tam=None, antes=0.0, desp=0.0):
    """Especificación de un nivel de título.
    alin: izq|centro · mayus: None|'caps'|'small' · sang: sangría izquierda (cm)."""
    return dict(alin=alin, b=b, i=i, mayus=mayus, sang=sang, tam=tam, antes=antes, desp=desp)


@dataclass(frozen=True)
class Norma:
    clave: str
    nombre: str
    fuente: str
    tam: float
    interlineado: float                 # 1.0 sencillo · 1.5 · 2.0 doble
    margenes: tuple                     # (superior, inferior, izquierdo, derecho) en cm
    sangria: float                      # sangría de primera línea (cm)
    alineacion: str                     # izq | just
    espacio_par: float                  # espacio posterior entre párrafos (pt)
    titulos: dict                       # nivel (1-5) -> T(...)
    num_lugar: str                      # encabezado | pie
    num_alin: str                       # izq | centro | der
    ref_sangria: float                  # sangría francesa de las referencias (cm)
    ref_interl: float
    ref_despues: float
    ref_tam: float
    ref_numeradas: str | None           # None | '[{n}]' | '{n}.'
    cita_bloque: float                  # sangría izquierda de citas en bloque (cm)
    familia_citas: str                  # autor-fecha | numerica | libre
    papel: str = "carta"                # carta | a4
    columnas: int = 1
    tam_tabla: float = 10
    tam_leyenda: float = 10
    leyenda_alin: str = "izq"
    tam_titulo: float | None = None     # título del documento (estilo "Title")
    encabezado_mayus: bool = False


NORMAS: dict[str, Norma] = {
    "apa7": Norma(
        "apa7", "APA 7.ª edición", "Times New Roman", 12, 2.0, (2.54, 2.54, 2.54, 2.54),
        1.27, "izq", 0,
        {1: T("centro", b=True), 2: T("izq", b=True), 3: T("izq", b=True, i=True),
         4: T("izq", b=True, sang=1.27), 5: T("izq", b=True, i=True, sang=1.27)},
        "encabezado", "der", 1.27, 2.0, 0, 12, None, 1.27, "autor-fecha",
        tam_tabla=11, tam_leyenda=12, encabezado_mayus=True),
    "ieee": Norma(
        "ieee", "IEEE", "Times New Roman", 10, 1.0, (1.9, 2.54, 1.59, 1.59),
        0.43, "just", 0,
        {1: T("centro", mayus="small", antes=12, desp=6), 2: T("izq", i=True, antes=6, desp=3),
         3: T("izq", i=True, sang=0.43, antes=3, desp=3), 4: T("izq", i=True, sang=0.86),
         5: T("izq", i=True, sang=0.86)},
        "pie", "centro", 0.9, 1.0, 0, 9, "[{n}]", 0.9, "numerica",
        tam_tabla=8, tam_leyenda=9, leyenda_alin="centro", tam_titulo=24),
    "mla9": Norma(
        "mla9", "MLA 9.ª edición", "Times New Roman", 12, 2.0, (2.54, 2.54, 2.54, 2.54),
        1.27, "izq", 0,
        {1: T("izq", b=True), 2: T("izq", i=True), 3: T("izq", b=True, i=True),
         4: T("izq", sang=1.27), 5: T("izq", i=True, sang=1.27)},
        "encabezado", "der", 1.27, 2.0, 0, 12, None, 1.27, "autor-fecha",
        tam_tabla=12, tam_leyenda=12),
    "chicago17": Norma(
        "chicago17", "Chicago 17.ª edición", "Times New Roman", 12, 2.0, (2.54, 2.54, 2.54, 2.54),
        1.27, "izq", 0,
        {1: T("centro", b=True), 2: T("centro"), 3: T("izq", b=True), 4: T("izq", i=True),
         5: T("izq")},
        "encabezado", "der", 1.27, 1.0, 12, 12, None, 1.27, "libre",
        tam_tabla=10, tam_leyenda=11),
    "vancouver": Norma(
        "vancouver", "Vancouver / ICMJE (biomédico)", "Arial", 12, 2.0, (2.5, 2.5, 2.5, 2.5),
        0.0, "izq", 0,
        {1: T("izq", b=True, mayus="caps"), 2: T("izq", b=True), 3: T("izq", b=True, i=True),
         4: T("izq", i=True), 5: T("izq", i=True)},
        "encabezado", "der", 0.0, 1.0, 6, 11, "{n}.", 1.27, "numerica", papel="a4",
        tam_tabla=10, tam_leyenda=11),
    "icontec": Norma(
        "icontec", "ICONTEC (NTC 1486, Colombia)", "Arial", 12, 1.5, (3.0, 3.0, 3.0, 2.0),
        0.0, "izq", 12,
        {1: T("izq", b=True, mayus="caps", desp=12), 2: T("izq", mayus="caps", desp=12),
         3: T("izq", b=True, desp=12), 4: T("izq", desp=12), 5: T("izq", i=True, desp=12)},
        "encabezado", "der", 1.27, 1.0, 12, 12, None, 1.27, "libre",
        tam_tabla=10, tam_leyenda=10),
}

PAPEL_CM = {"carta": (21.59, 27.94), "a4": (21.0, 29.7)}
ALIN = {"izq": WD_ALIGN_PARAGRAPH.LEFT, "centro": WD_ALIGN_PARAGRAPH.CENTER,
        "der": WD_ALIGN_PARAGRAPH.RIGHT, "just": WD_ALIGN_PARAGRAPH.JUSTIFY}


# ════════════════════════════════════════════════════════════════════════════
# 2. OPCIONES, INFORME Y ERRORES
# ════════════════════════════════════════════════════════════════════════════
class ErrorReformateo(Exception):
    """Error esperado con mensaje pensado para mostrarse al usuario."""


@dataclass
class Opciones:
    papel: str | None = None              # carta | a4 | original  (None = el de la norma)
    encabezado: str = ""                  # texto de encabezado corrido (APA profesional, etc.)
    apellido: str = ""                    # para el encabezado MLA ("Apellido 1")
    ordenar_refs: bool = False            # ordenar alfabéticamente la lista de referencias
    numerar_refs: bool = False            # anteponer [n] / n. si faltan (IEEE, Vancouver)
    detectar_titulos: bool = True         # detectar títulos sin estilo de título
    columnas: int | None = None           # 1 | 2 (None = el de la norma)
    exportar_pdf: bool = False            # generar también un PDF (requiere LibreOffice)


@dataclass
class Informe:
    entrada: str = ""
    salida: str = ""
    pdf: str = ""
    norma: str = ""
    avisos: list = field(default_factory=list)
    datos: dict = field(default_factory=dict)

    def texto(self) -> str:
        l = [f"Norma aplicada : {self.norma}", f"Archivo creado  : {self.salida}"]
        if self.pdf:
            l.append(f"PDF creado      : {self.pdf}")
        for k, v in self.datos.items():
            l.append(f"  · {k}: {v}")
        for a in self.avisos:
            l.append(f"  ⚠ {a}")
        return "\n".join(l)


# ════════════════════════════════════════════════════════════════════════════
# 3. UTILIDADES DE XML / WORD
# ════════════════════════════════════════════════════════════════════════════
def _norm(t: str) -> str:
    return re.sub(r"\s+", " ", t).strip()


def _texto_el(p_el) -> str:
    return "".join(t.text or "" for t in p_el.iter(qn("w:t")))


def _sin_acentos(s: str) -> str:
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def _runs(p: Paragraph):
    """Todos los runs del párrafo, incluidos los que están dentro de hipervínculos."""
    return [Run(r, p) for r in p._p.iter(qn("w:r"))]


def _limpiar_ppr(p_el, etiquetas):
    pPr = p_el.find(qn("w:pPr"))
    if pPr is None:
        return
    for e in etiquetas:
        for nodo in pPr.findall(qn(f"w:{e}")):
            pPr.remove(nodo)


def _fuente_rfonts(rPr, nombre):
    """Fija la fuente eliminando atributos de tema (asciiTheme…) que la anularían."""
    rf = rPr.get_or_add_rFonts()
    for a in list(rf.attrib):
        del rf.attrib[a]
    for a in ("ascii", "hAnsi", "cs", "eastAsia"):
        rf.set(qn(f"w:{a}"), nombre)


def _fmt_run(run: Run, fuente, tam, *, b=None, i=None, mayus="keep", negro=False):
    _fuente_rfonts(run._r.get_or_add_rPr(), fuente)
    run.font.size = Pt(tam)
    if b is not None:
        run.font.bold = b
    if i is not None:
        run.font.italic = i
    if mayus != "keep":
        run.font.all_caps = True if mayus == "caps" else False
        run.font.small_caps = True if mayus == "small" else False
    es_link = run.style is not None and "hyperlink" in (run.style.name or "").lower()
    if negro and not es_link:
        run.font.color.rgb = RGBColor(0, 0, 0)
    elif not es_link:  # elimina colores heredados de copiar/pegar; conserva los de hipervínculos
        rPr = run._r.rPr
        if rPr is not None:
            for c in rPr.findall(qn("w:color")):
                rPr.remove(c)


def _marca_parrafo(p: Paragraph, fuente, tam):
    """Tamaño/fuente de la marca de párrafo (afecta a párrafos vacíos)."""
    pPr = p._p.get_or_add_pPr()
    rPr = pPr.find(qn("w:rPr"))
    if rPr is None:
        rPr = OxmlElement("w:rPr")
        pPr.append(rPr)
    _fuente_rfonts(rPr, fuente)
    for tag in ("sz", "szCs"):
        e = rPr.find(qn(f"w:{tag}"))
        if e is None:
            e = OxmlElement(f"w:{tag}")
            rPr.append(e)
        e.set(qn("w:val"), str(int(tam * 2)))


def _fmt_par(p: Paragraph, *, alin="izq", primera=0.0, izq=0.0, colgante=0.0, interl=1.0,
             antes=0.0, despues=0.0, sig=False, tocar_sangria=True, tocar_alin=True):
    """Aplica formato de párrafo limpiando antes el formato directo previo."""
    borrar = ["spacing", "contextualSpacing", "keepNext", "keepLines", "widowControl"]
    if tocar_sangria:
        borrar += ["ind", "mirrorIndents"]
    if tocar_alin:
        borrar += ["jc"]
    _limpiar_ppr(p._p, borrar)
    pf = p.paragraph_format
    if tocar_alin:
        pf.alignment = ALIN[alin]
    if tocar_sangria:
        pf.left_indent = Cm(izq + colgante)
        pf.first_line_indent = Cm(-colgante) if colgante else Cm(primera)
    pf.space_before = Pt(antes)
    pf.space_after = Pt(despues)
    pf.line_spacing = interl
    pf.widow_control = True
    pf.keep_with_next = True if sig else None


def _estilo_fuente(doc, nombre_estilo, fuente, tam):
    try:
        st = doc.styles[nombre_estilo]
    except KeyError:
        return
    _fuente_rfonts(st.element.get_or_add_rPr(), fuente)
    st.font.size = Pt(tam)


def _estilo_titulo(doc, nivel):
    nombre = f"Heading {nivel}"
    try:
        return doc.styles[nombre]
    except KeyError:
        st = doc.styles.add_style(nombre, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = doc.styles["Normal"]
        ppr = st.element.get_or_add_pPr()
        ol = OxmlElement("w:outlineLvl")
        ol.set(qn("w:val"), str(nivel - 1))
        ppr.append(ol)
        return st


def _campo(par: Paragraph, instr: str, fuente, tam):
    """Inserta un campo de Word (p. ej. PAGE) con la fuente indicada."""
    def nuevo_run():
        r = par.add_run()
        _fmt_run(r, fuente, tam)
        return r
    r = nuevo_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "begin"); r._r.append(e)
    r = nuevo_run(); e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = f" {instr} "; r._r.append(e)
    r = nuevo_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "separate"); r._r.append(e)
    r = nuevo_run(); r.text = "1"
    r = nuevo_run(); e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end"); r._r.append(e)


# ════════════════════════════════════════════════════════════════════════════
# 4. LECTURA DE PDF  →  DOCX NEUTRO
# ════════════════════════════════════════════════════════════════════════════
RE_NUM_PAG = re.compile(r"^\s*(?:p[áa]g(?:ina)?\.?|page)?\s*\d+(?:\s*(?:de|of|/)\s*\d+)?\s*$", re.I)


def pdf_a_docx(ruta_pdf: Path, ruta_docx: Path, avisos: list):
    """Extrae el texto del PDF (con negritas/cursivas, títulos y párrafos) y lo guarda
    como un .docx sencillo que luego pasa por el mismo formateador que un .docx."""
    try:
        import pdfplumber
    except ImportError:
        raise ErrorReformateo("Para leer PDF se necesita 'pdfplumber':  pip install pdfplumber")

    lineas = []
    n_img = 0
    try:
        with pdfplumber.open(str(ruta_pdf)) as pdf:
            n_pag = len(pdf.pages)
            for num, pag in enumerate(pdf.pages):
                n_img += len(pag.images)
                palabras = pag.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False,
                                             extra_attrs=["fontname", "size"])
                palabras.sort(key=lambda w: (w["top"], w["x0"]))
                grupo, ref = [], None

                def cerrar():
                    if not grupo:
                        return
                    g = sorted(grupo, key=lambda w: w["x0"])
                    toks = []
                    for w in g:
                        f = w["fontname"].lower()
                        toks.append((w["text"], any(k in f for k in ("bold", "black", "heavy")),
                                     "italic" in f or "oblique" in f))
                    hueco = max((b["x0"] - a["x1"] for a, b in zip(g, g[1:])), default=0)
                    lineas.append(dict(
                        pag=num, alto=pag.height, ancho=pag.width, top=min(w["top"] for w in g),
                        bottom=max(w["bottom"] for w in g), x0=g[0]["x0"], x1=max(w["x1"] for w in g),
                        size=statistics.median(w["size"] for w in g), toks=toks,
                        texto=" ".join(t[0] for t in toks), negrita=all(t[1] for t in toks),
                        hueco=hueco))

                for w in palabras:
                    if grupo and abs(w["top"] - ref) <= max(2.0, 0.4 * w["size"]):
                        grupo.append(w)
                    else:
                        cerrar(); grupo = [w]; ref = w["top"]
                cerrar()
    except ErrorReformateo:
        raise
    except Exception as e:
        if "password" in str(e).lower() or "encrypt" in str(e).lower():
            raise ErrorReformateo("El PDF está protegido con contraseña. Quite la protección e inténtelo de nuevo.")
        raise ErrorReformateo(f"No se pudo leer el PDF ({type(e).__name__}: {e}).")

    if not lineas:
        raise ErrorReformateo("El PDF no contiene texto seleccionable (probablemente es un escaneo). "
                              "Aplique OCR primero (p. ej. con ocrmypdf) y vuelva a intentarlo.")

    # — eliminar encabezados, pies y números de página —
    def clave(l):
        return re.sub(r"\d+", "#", l["texto"].strip().lower())
    conteo = collections.Counter()
    for k in {(clave(l), l["pag"]) for l in lineas}:
        conteo[k[0]] += 1
    filtradas = []
    for l in lineas:
        banda = l["top"] < 0.08 * l["alto"] or l["bottom"] > 0.92 * l["alto"]
        repetida = n_pag >= 2 and conteo[clave(l)] >= max(2, 0.4 * n_pag)
        if banda and (RE_NUM_PAG.match(l["texto"]) or repetida):
            continue
        filtradas.append(l)
    lineas = filtradas or lineas

    # — tamaño de cuerpo y márgenes —
    pesos = collections.Counter()
    for l in lineas:
        pesos[round(l["size"] * 2) / 2] += len(l["texto"])
    cuerpo = pesos.most_common(1)[0][0]
    izq = statistics.median_low([l["x0"] for l in lineas])
    der = sorted(l["x1"] for l in lineas)[int(0.95 * (len(lineas) - 1))]
    if sum(1 for l in lineas if l["hueco"] > 0.2 * l["ancho"]) > 0.15 * len(lineas):
        avisos.append("El PDF parece tener varias columnas: el orden de lectura puede no ser el correcto. "
                      "Si tiene el original en Word, úselo para un resultado más fiable.")

    # — separar líneas de título —
    n = len(lineas)
    def gap(i, j):  # distancia vertical entre líneas de la misma página
        a, b = lineas[i], lineas[j]
        return b["top"] - a["bottom"] if a["pag"] == b["pag"] else 999
    tit = [None] * n
    for i, l in enumerate(lineas):
        if l["size"] >= cuerpo * 1.12:
            tit[i] = ("size", round(l["size"]))
        elif (l["negrita"] and 2 < len(l["texto"]) <= 100 and not l["texto"].endswith(".")
              and (i == 0 or gap(i - 1, i) > 0.5 * cuerpo) and (i == n - 1 or gap(i, i + 1) > 0.5 * cuerpo)):
            tit[i] = ("negrita", 0)
    tamanos = sorted({t[1] for t in tit if t and t[0] == "size"}, reverse=True)
    nivel_de = {s: min(k + 1, 5) for k, s in enumerate(tamanos)}
    nivel_neg = min(len(tamanos) + 1, 5)

    # — construir bloques —
    bloques, actual, prev = [], None, None  # bloque: dict(tipo, nivel, palabras)
    for i, l in enumerate(lineas):
        if tit[i]:
            niv = nivel_de.get(tit[i][1], nivel_neg) if tit[i][0] == "size" else nivel_neg
            if actual and actual["tipo"] == "t" and actual["nivel"] == niv and prev is not None \
                    and gap(prev, i) < 1.2 * l["size"]:
                actual["toks"] += l["toks"]
            else:
                actual = dict(tipo="t", nivel=niv, toks=list(l["toks"])); bloques.append(actual)
        else:
            nuevo = actual is None or actual["tipo"] == "t"
            if not nuevo:
                a = lineas[prev]
                if a["pag"] == l["pag"]:
                    g = gap(prev, i)
                    corta = a["x1"] < der - 0.12 * (der - izq) and a["texto"].rstrip()[-1:] in ".:;?!"
                    nuevo = g > 0.6 * cuerpo or l["x0"] > izq + 1.2 * cuerpo or corta
                else:
                    nuevo = a["texto"].rstrip()[-1:] in ".:;?!" and l["texto"][:1].isupper()
            if nuevo:
                actual = dict(tipo="p", nivel=0, toks=[]); bloques.append(actual)
            toks = list(l["toks"])
            if actual["toks"] and toks:  # deshacer guiones de fin de línea (palabra-\ncontinuación)
                t0, (t1, b1, i1) = toks[0], actual["toks"][-1]
                if t1.endswith("-") and len(t1) > 2 and t1[-2].isalpha() and t0[0][:1].islower():
                    actual["toks"][-1] = (t1[:-1] + t0[0], b1, i1)
                    toks = toks[1:]
            actual["toks"] += toks
        prev = i

    # — escribir el docx neutro —
    doc = Document()
    for b in bloques:
        if b["tipo"] == "t":
            p = doc.add_paragraph(style=f"Heading {b['nivel']}")
            p.add_run(" ".join(t[0] for t in b["toks"]))
        else:
            p = doc.add_paragraph()
            runs = []
            for texto, neg, cur in b["toks"]:
                if runs and runs[-1][1:] == [neg, cur]:
                    runs[-1][0] += " " + texto
                else:
                    if runs:
                        runs[-1][0] += " "
                    runs.append([texto, neg, cur])
            for texto, neg, cur in runs:
                r = p.add_run(texto)
                r.bold = True if neg else None
                r.italic = True if cur else None
    doc.save(str(ruta_docx))
    if n_img:
        avisos.append(f"El PDF contiene {n_img} imagen(es) y posibles tablas: desde un PDF solo se recupera el "
                      "texto. Para conservar imágenes y tablas use el documento original en Word.")
    avisos.append("Origen PDF: revise los títulos detectados y los saltos de párrafo en el resultado.")


# ════════════════════════════════════════════════════════════════════════════
# 5. CLASIFICACIÓN DE PÁRRAFOS (.docx)
# ════════════════════════════════════════════════════════════════════════════
RE_HEAD_ESTILO = re.compile(r"^(?:heading|t[ií]tulo|titulo)\s*(\d)$", re.I)
RE_TITLE_ESTILO = re.compile(r"^(?:title|t[ií]tulo|titulo)$", re.I)
RE_NUM_TIT = re.compile(r"^(\d+(?:\.\d+){0,4})[\.\)]?\s+\S")
RE_ROMANO = re.compile(r"^[IVX]+\.\s+\S")
RE_LEYENDA = re.compile(r"^(?:tabla|table|figura|figure|fig\.|gr[áa]fico|gr[áa]fica|ilustraci[óo]n|imagen|cuadro)\s+\d+", re.I)
RE_LISTA = re.compile(r"^\s*(?:[•●▪◦■\-–—*]|\(?\d{1,2}[\.\)]|\(?[a-z]\))\s+\S")
RE_REFS = re.compile(r"^(?:\d+(?:\.\d+)*[\.\)]?\s*)?(?:lista\s+de\s+)?(?:referencias(?:\s+bibliogr[áa]ficas)?|bibliograf[íi]a|references|bibliography|works\s+cited|obras\s+citadas|fuentes\s+consultadas)\s*$", re.I)
RE_ANEXOS = re.compile(r"^(?:\d+[\.\)]?\s*)?(?:anexos?|ap[ée]ndices?|appendix|appendices)\b", re.I)
RE_YA_NUM = re.compile(r"^\s*(?:\[\d+\]|\(\d+\)|\d+[\.\)])\s")
SECCIONES = {"resumen", "abstract", "introduccion", "conclusiones", "conclusion", "recomendaciones",
             "referencias", "bibliografia", "agradecimientos", "anexos", "apendices", "metodologia",
             "resultados", "discusion", "marco teorico", "objetivos", "justificacion", "dedicatoria",
             "planteamiento del problema", "palabras clave", "keywords", "tabla de contenido", "indice",
             "contenido", "lista de tablas", "lista de figuras", "references", "introduction",
             "methods", "results", "discussion", "conclusions", "acknowledgments"}
RE_CIT_AUTOR = re.compile(r"\((?:[A-ZÁÉÍÓÚÑ][^()]{1,80}?),?\s+(?:\d{4}[a-z]?|s\.\s?f\.|n\.\s?d\.)(?:,\s*p+\.\s*\d+[^()]*)?\)")
RE_CIT_NUM = re.compile(r"\[\d+(?:\s*[,–\-]\s*\d+)*\]")


@dataclass
class Info:
    p: Paragraph
    kind: str = "cuerpo"     # cuerpo|titulo_doc|subtitulo|heading|toc|toc_titulo|leyenda|imagen|lista|
    #                          lista_manual|cita|ref|tabla|omitir|vacio
    nivel: int = 0
    texto: str = ""
    centrado: bool = False


def _estilo_nombre(p: Paragraph) -> str:
    try:
        return (p.style.name or "").strip() if p.style is not None else ""
    except Exception:
        return ""


def _es_negrita(p: Paragraph) -> bool:
    runs = [r for r in _runs(p) if (r.text or "").strip()]
    if not runs:
        return False
    st_b = bool(p.style is not None and p.style.font.bold)
    return all((r.bold if r.bold is not None else st_b) for r in runs)


def clasificar(doc) -> list[Info]:
    body = doc.element.body
    infos = []
    for p_el in body.iter(qn("w:p")):
        p = Paragraph(p_el, doc)
        txt = _texto_el(p_el)
        inf = Info(p=p, texto=txt)
        est = _estilo_nombre(p)
        est_l = est.lower()
        if p_el.xpath("ancestor::w:txbxContent"):
            inf.kind = "omitir"
        elif p_el.xpath("ancestor::w:tbl"):
            inf.kind = "tabla"
        elif est_l.startswith("toc heading"):
            inf.kind, inf.nivel = "toc_titulo", 1
        elif est_l.startswith(("toc", "table of contents", "tdc", "índice", "indice", "tabla de contenido")) \
                or est_l.startswith("table of figures"):
            inf.kind = "toc"
        elif (m := RE_HEAD_ESTILO.match(est)):
            inf.kind, inf.nivel = "heading", max(1, min(int(m.group(1)), 5))
        elif RE_TITLE_ESTILO.match(est):
            inf.kind = "titulo_doc"
        elif est_l in ("subtitle", "subtítulo", "subtitulo"):
            inf.kind = "subtitulo"
        elif not txt.strip():
            inf.kind = "imagen" if p_el.xpath(".//w:drawing|.//w:pict") else "vacio"
        elif p_el.xpath(".//w:drawing|.//w:pict") and len(txt.strip()) < 3:
            inf.kind = "imagen"
        elif est_l in ("caption", "epígrafe", "epigrafe", "leyenda") or (RE_LEYENDA.match(txt.strip()) and len(txt) < 250):
            inf.kind = "leyenda"
        elif est_l in ("bibliography", "bibliografía", "bibliografia"):
            inf.kind = "ref"
        elif est_l in ("quote", "intense quote", "block text", "cita", "cita destacada", "texto de bloque"):
            inf.kind = "cita"
        elif p_el.find(qn("w:pPr")) is not None and p_el.find(qn("w:pPr")).find(qn("w:numPr")) is not None:
            inf.kind = "lista"
        elif "list" in est_l or "lista" in est_l:
            inf.kind = "lista"
        inf.centrado = p.alignment in (WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT) and len(txt) <= 150
        infos.append(inf)
    return infos


def detectar_titulos_heuristicos(infos: list[Info]) -> int:
    """Marca como título los párrafos cortos que lo parecen (sin estilo de título)."""
    n = 0
    for inf in infos:
        if inf.kind != "cuerpo":
            continue
        t = inf.texto.strip()
        if not t or len(t) > 120 or len(t.split()) > 18 or t.endswith((".", ",", ";", ":")):
            continue
        clave = _sin_acentos(re.sub(r"^[\dIVX]+[\.\)]?\s*(?:\d+[\.\)]?\s*)*", "", t)).lower().strip(" .:")
        niv = 0
        if clave in SECCIONES:
            niv = 1
        elif (m := RE_NUM_TIT.match(t)) and len(t) <= 90:
            niv = min(m.group(1).count(".") + 1, 5)
        elif RE_ROMANO.match(t) and len(t) <= 90:
            niv = 1
        elif _es_negrita(inf.p):
            niv = 1 if (inf.centrado or t.isupper()) else 2
        elif t.isupper() and 3 < len(t) <= 80:
            niv = 1
        if niv:
            inf.kind, inf.nivel = "heading", niv
            n += 1
    return n


def marcar_lista_y_citas(infos: list[Info]):
    for inf in infos:
        if inf.kind != "cuerpo":
            continue
        t = inf.texto.strip()
        if RE_LISTA.match(inf.texto):
            inf.kind = "lista_manual"
        elif len(t.split()) >= 40 and t[0] in "\"“«" and t[-1] in "\"”»":
            inf.kind = "cita"


def marcar_referencias(infos: list[Info]) -> list[Info]:
    """Marca como 'ref' los párrafos de la sección de referencias. Devuelve esos párrafos."""
    refs, dentro, nivel_ref = [], False, 0
    for inf in infos:
        t = inf.texto.strip()
        if inf.kind == "heading":
            if RE_REFS.match(t):
                dentro, nivel_ref = True, inf.nivel
                continue
            if dentro and inf.nivel <= nivel_ref:
                dentro = False
        elif dentro and (RE_ANEXOS.match(t) and len(t) < 60):
            dentro = False
        if dentro and inf.kind in ("cuerpo", "lista", "lista_manual", "ref", "vacio") :
            if inf.kind != "vacio":
                inf.kind = "ref"
                refs.append(inf)
    # un párrafo con estilo Bibliography fuera de la sección también cuenta
    for inf in infos:
        if inf.kind == "ref" and inf not in refs:
            refs.append(inf)
    return refs


def ordenar_referencias(refs: list[Info]) -> bool:
    """Reordena alfabéticamente los párrafos de referencias (solo mueve párrafos, no cambia texto)."""
    if len(refs) < 2:
        return False
    padres = {id(i.p._p.getparent()) for i in refs}
    if len(padres) != 1:
        return False
    def k(inf):
        t = re.sub(r"^\s*(?:\[\d+\]|\(\d+\)|\d+[\.\)])\s*", "", inf.texto)
        return _sin_acentos(t).lstrip(" \"“«([").casefold()
    padre = refs[0].p._p.getparent()
    hijos = list(padre)
    i0 = hijos.index(refs[0].p._p)
    i1 = hijos.index(refs[-1].p._p)
    bloque = hijos[i0:i1 + 1]
    if any(e.tag not in (qn("w:p"),) for e in bloque):
        return False
    ref_els = {i.p._p for i in refs}
    ordenados = iter(sorted([i for i in refs], key=k))
    nuevo = [next(ordenados).p._p if e in ref_els else e for e in bloque]
    if nuevo == bloque:
        return False
    for e in bloque:
        padre.remove(e)
    for j, e in enumerate(nuevo):
        padre.insert(i0 + j, e)
    return True


# ════════════════════════════════════════════════════════════════════════════
# 6. APLICACIÓN DE FORMATO
# ════════════════════════════════════════════════════════════════════════════
def configurar_paginas(doc, n: Norma, op: Opciones):
    papel = op.papel or n.papel
    cols = op.columnas or n.columnas
    sup, inf, izq, der = n.margenes
    for s in doc.sections:
        s.top_margin, s.bottom_margin = Cm(sup), Cm(inf)
        s.left_margin, s.right_margin = Cm(izq), Cm(der)
        s.gutter = Cm(0)
        s.header_distance = Cm(min(1.25, sup - 0.5))
        s.footer_distance = Cm(min(1.25, inf - 0.5))
        if papel in PAPEL_CM:
            w, h = PAPEL_CM[papel]
            if s.orientation == WD_ORIENT.LANDSCAPE:
                w, h = h, w
            s.page_width, s.page_height = Cm(w), Cm(h)
        pg = s._sectPr
        c = pg.find(qn("w:cols"))
        if cols > 1 or c is not None:
            if c is None:
                c = OxmlElement("w:cols")
                pg.find(qn("w:pgMar")).addnext(c)
            c.set(qn("w:num"), str(cols))
            if cols > 1:
                c.set(qn("w:space"), str(int(0.63 / 2.54 * 1440)))  # 0,63 cm entre columnas


def aplicar_formato(doc, n: Norma, infos: list[Info], op: Opciones) -> dict:
    alin_cuerpo = n.alineacion
    stats = collections.Counter()
    _estilo_fuente(doc, "Normal", n.fuente, n.tam)

    for inf in infos:
        p, k = inf.p, inf.kind
        if k == "omitir":
            continue
        stats[k] += 1
        fuente, tam = n.fuente, n.tam

        if k == "heading":
            sp = n.titulos[inf.nivel]
            if not p.style.name.lower().startswith(("heading", "título", "titulo")):
                p.style = _estilo_titulo(doc, inf.nivel)
            tam = sp["tam"] or n.tam
            prev = p._p.getprevious()
            tras_tabla = prev is not None and prev.tag == qn("w:tbl")  # evita títulos pegados a una tabla
            _fmt_par(p, alin=sp["alin"], izq=sp["sang"], interl=n.interlineado if n.interlineado > 1 else 1.0,
                     antes=max(sp["antes"], 12 if tras_tabla else 0), despues=sp["desp"], sig=True)
            for r in _runs(p):
                _fmt_run(r, fuente, tam, b=sp["b"], i=sp["i"], mayus=sp["mayus"], negro=True)
            _marca_parrafo(p, fuente, tam)

        elif k in ("titulo_doc", "subtitulo"):
            tam = n.tam_titulo or n.tam
            _fmt_par(p, alin="centro", interl=n.interlineado if n.interlineado > 1 else 1.0, despues=6, sig=True)
            for r in _runs(p):
                _fmt_run(r, fuente, tam, b=(k == "titulo_doc"), negro=True)
            _marca_parrafo(p, fuente, tam)

        elif k in ("toc", "toc_titulo"):
            # Solo fuente e interlineado; sangrías y tabuladores del índice se respetan.
            _fmt_par(p, interl=1.0, despues=0, tocar_sangria=False, tocar_alin=(k == "toc_titulo"),
                     alin="centro")
            for r in _runs(p):
                _fmt_run(r, fuente, tam, negro=(k == "toc_titulo"))

        elif k == "tabla":
            tam = n.tam_tabla
            _fmt_par(p, interl=1.0, despues=0, tocar_sangria=False, tocar_alin=False)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        elif k == "leyenda":
            tam = n.tam_leyenda
            sig = inf.texto.strip().lower().startswith(("tabla", "table", "cuadro"))
            _fmt_par(p, alin=n.leyenda_alin, interl=1.0, antes=6 if sig else 0, despues=6, sig=sig)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        elif k == "imagen":
            _fmt_par(p, alin="centro", interl=1.0, despues=6, sig=True)
            _marca_parrafo(p, fuente, tam)

        elif k == "lista":  # numeración de Word: se respeta su sangría
            _fmt_par(p, alin=alin_cuerpo, interl=n.interlineado, despues=n.espacio_par,
                     tocar_sangria=False)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        elif k == "lista_manual":
            _fmt_par(p, alin=alin_cuerpo, izq=1.27, colgante=0.63, interl=n.interlineado, despues=n.espacio_par)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        elif k == "cita":
            _fmt_par(p, alin="izq" if alin_cuerpo != "just" else "just", izq=n.cita_bloque,
                     interl=n.interlineado, despues=n.espacio_par)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        elif k == "ref":
            tam = n.ref_tam
            _fmt_par(p, alin="izq" if alin_cuerpo != "just" else "just", izq=0, colgante=n.ref_sangria,
                     interl=n.ref_interl, despues=n.ref_despues)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)

        else:  # cuerpo / vacío
            orig = p.alignment
            if k == "cuerpo" and inf.centrado:   # portadas, firmas, fechas: se respeta el centrado/derecha
                _fmt_par(p, alin="centro" if orig == WD_ALIGN_PARAGRAPH.CENTER else "der",
                         interl=n.interlineado, despues=n.espacio_par)
            else:
                _fmt_par(p, alin=alin_cuerpo, primera=n.sangria if k == "cuerpo" else 0,
                         interl=n.interlineado, despues=n.espacio_par)
            for r in _runs(p):
                _fmt_run(r, fuente, tam)
            _marca_parrafo(p, fuente, tam)
    return dict(stats)


def numerar_referencias(doc, n: Norma, refs: list[Info]) -> int:
    if not n.ref_numeradas:
        return 0
    cuenta = 0
    for k, inf in enumerate(refs, start=1):
        if RE_YA_NUM.match(inf.texto):
            continue
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = n.ref_numeradas.format(n=k) + " "
        r.append(t)
        pPr = inf.p._p.find(qn("w:pPr"))
        if pPr is not None:
            pPr.addnext(r)
        else:
            inf.p._p.insert(0, r)
        _fmt_run(Run(r, inf.p), n.fuente, n.ref_tam)
        cuenta += 1
    return cuenta


def configurar_numeracion(doc, n: Norma, op: Opciones):
    """Numeración de páginas (campo PAGE) y encabezado corrido según la norma."""
    ajustes = doc.settings.element
    for e in ajustes.findall(qn("w:evenAndOddHeaders")):
        ajustes.remove(e)
    _estilo_fuente(doc, "Header", n.fuente, n.tam)
    _estilo_fuente(doc, "Footer", n.fuente, n.tam)
    tam = n.tam
    prefijo = (op.apellido.strip() + " ") if (n.clave == "mla9" and op.apellido.strip()) else ""
    for s in doc.sections:
        s.different_first_page_header_footer = False
        for hf in (s.header, s.footer):
            hf.is_linked_to_previous = False
            pars = hf.paragraphs
            for extra in pars[1:]:
                extra._p.getparent().remove(extra._p)
            for hijo in list(pars[0]._p):
                if hijo.tag != qn("w:pPr"):
                    pars[0]._p.remove(hijo)
        destino = s.header if n.num_lugar == "encabezado" else s.footer
        par = destino.paragraphs[0]
        pf = par.paragraph_format
        pf.space_before = pf.space_after = Pt(0)
        pf.line_spacing = 1.0
        ancho = s.page_width - s.left_margin - s.right_margin
        # quitar tabuladores heredados del estilo Header/Footer
        for ts in list(par.style.paragraph_format.tab_stops):
            pf.tab_stops.add_tab_stop(ts.position, WD_TAB_ALIGNMENT.CLEAR)
        if op.encabezado.strip() and n.num_lugar == "encabezado":
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.tab_stops.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT)
            r = par.add_run(op.encabezado.strip() + "\t")
            _fmt_run(r, n.fuente, tam, mayus="caps" if n.encabezado_mayus else "keep")
        else:
            pf.alignment = ALIN[n.num_alin]
            if prefijo:
                _fmt_run(par.add_run(prefijo), n.fuente, tam)
        _campo(par, "PAGE", n.fuente, tam)


def forzar_actualizar_campos(doc) -> bool:
    """Si hay tabla de contenido, pide a Word que actualice campos al abrir."""
    body = doc.element.body
    hay = any("TOC" in (e.text or "") for e in body.iter(qn("w:instrText"))) or \
        any("TOC" in (e.get(qn("w:instr")) or "") for e in body.iter(qn("w:fldSimple")))
    if hay:
        u = OxmlElement("w:updateFields")
        u.set(qn("w:val"), "true")
        doc.settings.element.append(u)
    return hay


def verificar_texto(doc_entrada_textos: collections.Counter, doc) -> bool:
    final = collections.Counter(t for t in (_norm(_texto_el(p)) for p in doc.element.body.iter(qn("w:p"))) if t)
    return final == doc_entrada_textos


def _textos(doc) -> collections.Counter:
    return collections.Counter(t for t in (_norm(_texto_el(p)) for p in doc.element.body.iter(qn("w:p"))) if t)


# ════════════════════════════════════════════════════════════════════════════
# 7. CONVERSIONES CON LIBREOFFICE (opcionales)
# ════════════════════════════════════════════════════════════════════════════
def _soffice():
    return shutil.which("soffice") or shutil.which("libreoffice") or next(
        (p for p in (r"C:\Program Files\LibreOffice\program\soffice.exe",
                     r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
                     "/Applications/LibreOffice.app/Contents/MacOS/soffice") if Path(p).exists()), None)


def _convertir(ruta: Path, formato: str, carpeta: Path) -> Path:
    exe = _soffice()
    if not exe:
        return None
    try:
        subprocess.run([exe, "--headless", "--convert-to", formato, "--outdir", str(carpeta), str(ruta)],
                       check=True, capture_output=True, timeout=180)
    except Exception:
        return None
    res = carpeta / (ruta.stem + "." + formato)
    return res if res.exists() else None


# ════════════════════════════════════════════════════════════════════════════
# 8. FLUJO PRINCIPAL
# ════════════════════════════════════════════════════════════════════════════
def reformatear(entrada, norma: str = "apa7", salida=None, opciones: Opciones | None = None,
                log=lambda m: None) -> Informe:
    op = opciones or Opciones()
    if norma not in NORMAS:
        raise ErrorReformateo(f"Norma desconocida '{norma}'. Opciones: {', '.join(NORMAS)}")
    n = NORMAS[norma]
    ruta = Path(str(entrada)).expanduser()
    if not ruta.is_file():
        raise ErrorReformateo(f"No se encontró el archivo: {ruta}")
    ext = ruta.suffix.lower()
    if ext not in (".docx", ".pdf", ".doc", ".odt", ".rtf", ".docm"):
        raise ErrorReformateo(f"Formato '{ext}' no soportado. Use .docx o .pdf (también .doc/.odt/.rtf con LibreOffice).")

    inf = Informe(entrada=str(ruta), norma=n.nombre)
    salida = Path(salida) if salida else ruta.with_name(f"{ruta.stem}_{n.clave}.docx")
    if salida.suffix.lower() != ".docx":
        salida = salida.with_suffix(".docx")
    if salida.resolve() == ruta.resolve():
        raise ErrorReformateo("El archivo de salida no puede ser el mismo que el de entrada.")

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        fuente = ruta
        if ext == ".pdf":
            log("Leyendo PDF y reconstruyendo títulos y párrafos…")
            fuente = tmp / "desde_pdf.docx"
            pdf_a_docx(ruta, fuente, inf.avisos)
        elif ext != ".docx":
            log(f"Convirtiendo {ext} a .docx con LibreOffice…")
            conv = _convertir(ruta, "docx", tmp)
            if not conv:
                raise ErrorReformateo(f"Para abrir archivos {ext} se necesita LibreOffice instalado. "
                                      "Guárdelo como .docx desde Word y vuelva a intentarlo.")
            fuente = conv

        log("Abriendo documento…")
        try:
            doc = Document(str(fuente))
        except Exception as e:
            raise ErrorReformateo("No se pudo abrir el .docx: está dañado, protegido con contraseña o no es un "
                                  f"archivo de Word válido ({type(e).__name__}).")

        textos0 = _textos(doc)
        if not textos0:
            raise ErrorReformateo("El documento no contiene texto.")

        log("Analizando estructura…")
        infos = clasificar(doc)
        n_estilos = sum(1 for i in infos if i.kind == "heading")
        if op.detectar_titulos and n_estilos < 2:
            det = detectar_titulos_heuristicos(infos)
            if det:
                inf.avisos.append(f"El documento casi no usa estilos de título; se detectaron {det} título(s) "
                                  "por su aspecto (revíselos).")
        marcar_lista_y_citas(infos)
        refs = marcar_referencias(infos)
        if not refs:
            inf.avisos.append("No se encontró una sección de 'Referencias/Bibliografía'; no se aplicó sangría francesa.")
        elif op.ordenar_refs:
            if ordenar_referencias(refs):
                refs = [i for i in infos if i.kind == "ref"]
                log("Referencias ordenadas alfabéticamente.")

        log(f"Aplicando formato {n.nombre}…")
        configurar_paginas(doc, n, op)
        stats = aplicar_formato(doc, n, infos, op)

        log("Verificando que el texto no cambió…")
        if not verificar_texto(textos0, doc):
            inf.avisos.append("¡ATENCIÓN! El texto de salida difiere del original. Revise el resultado.")
        else:
            inf.datos["Texto verificado"] = "idéntico al original"

        if op.numerar_refs:
            if n.ref_numeradas:
                k = numerar_referencias(doc, n, refs)
                inf.datos["Referencias numeradas"] = k
                inf.avisos.append("Se añadieron números a las referencias: las citas dentro del texto deben "
                                  "usar esos mismos números.")
            else:
                inf.avisos.append(f"{n.nombre} no usa referencias numeradas; se ignoró 'numerar referencias'.")

        configurar_numeracion(doc, n, op)
        if forzar_actualizar_campos(doc):
            inf.avisos.append("Hay tabla de contenido: al abrir en Word acepte 'actualizar campos' para "
                              "refrescar los números de página.")

        # — comprobación de estilo de citas dentro del texto —
        cuerpo = " ".join(i.texto for i in infos if i.kind in ("cuerpo", "lista", "cita"))
        c_af, c_num = len(RE_CIT_AUTOR.findall(cuerpo)), len(RE_CIT_NUM.findall(cuerpo))
        inf.datos["Citas detectadas"] = f"{c_af} autor-fecha, {c_num} numéricas"
        if n.familia_citas == "numerica" and c_af > c_num:
            inf.avisos.append(f"{n.nombre} usa citas numéricas [1], pero el texto usa autor-fecha (Autor, año). "
                              "El script NO convierte citas: debe reescribirlas manualmente.")
        elif n.familia_citas == "autor-fecha" and c_num > c_af:
            inf.avisos.append(f"{n.nombre} usa citas autor-fecha, pero el texto usa citas numéricas. "
                              "El script NO convierte citas: debe reescribirlas manualmente.")
        inf.avisos.append("No se modificaron los bordes de las tablas ni el texto de las citas/referencias "
                          "(solo su formato).")

        inf.datos["Elementos"] = ", ".join(f"{v} {k}" for k, v in sorted(stats.items()) if k not in ("omitir",))
        log("Guardando…")
        try:
            salida.parent.mkdir(parents=True, exist_ok=True)
            doc.save(str(salida))
        except PermissionError:
            raise ErrorReformateo(f"No se pudo guardar '{salida.name}': ciérrelo si lo tiene abierto en Word.")
        inf.salida = str(salida)

        if op.exportar_pdf:
            log("Generando PDF…")
            pdf = _convertir(salida, "pdf", salida.parent)
            if pdf:
                inf.pdf = str(pdf)
            else:
                inf.avisos.append("No se pudo generar el PDF (requiere LibreOffice instalado).")
    return inf


# ════════════════════════════════════════════════════════════════════════════
# 9. INTERFAZ GRÁFICA (arrastrar y soltar)
# ════════════════════════════════════════════════════════════════════════════
def abrir_carpeta(ruta: str):
    import os
    carpeta = str(Path(ruta).parent)
    try:
        if sys.platform.startswith("win"):
            os.startfile(carpeta)  # type: ignore[attr-defined]
        elif sys.platform == "darwin":
            subprocess.Popen(["open", carpeta])
        else:
            subprocess.Popen(["xdg-open", carpeta])
    except Exception:
        pass


def lanzar_gui():
    try:
        import tkinter as tk
        from tkinter import filedialog, messagebox, ttk
    except ImportError:
        sys.exit("Tkinter no está disponible en este Python. Use la línea de comandos: "
                 "python reformateador.py archivo.docx -n apa7")
    try:
        from tkinterdnd2 import DND_FILES, TkinterDnD
        root = TkinterDnD.Tk()
        dnd = True
    except Exception:
        root = tk.Tk()
        dnd = False

    root.title("Reformateador de documentos académicos")
    root.geometry("680x640")
    root.minsize(600, 560)
    estado = {"archivo": None}
    cola: queue.Queue = queue.Queue()

    marco = ttk.Frame(root, padding=14)
    marco.pack(fill="both", expand=True)

    zona = tk.Label(marco, text="", relief="groove", borderwidth=3, height=6, bg="#f3f6fb", fg="#33415c",
                    font=("Segoe UI", 11), cursor="hand2", wraplength=560, justify="center")
    zona.pack(fill="x", pady=(0, 10))

    def texto_zona():
        if estado["archivo"]:
            zona.config(text=f"📄  {Path(estado['archivo']).name}\n\n(arrastre otro archivo para cambiarlo)")
        else:
            extra = "" if dnd else "\n\n(Arrastrar y soltar requiere:  pip install tkinterdnd2)"
            zona.config(text="⬇  Arrastre aquí su documento (.docx o .pdf)\no haga clic para buscarlo" + extra)

    def elegir(ruta):
        ruta = ruta.strip().strip("{}")
        if not ruta:
            return
        if Path(ruta).suffix.lower() not in (".docx", ".pdf", ".doc", ".odt", ".rtf", ".docm"):
            messagebox.showwarning("Formato no soportado", "Use un archivo .docx o .pdf")
            return
        estado["archivo"] = ruta
        texto_zona()

    def buscar(_=None):
        r = filedialog.askopenfilename(filetypes=[("Documentos", "*.docx *.pdf *.doc *.odt *.rtf"), ("Todos", "*.*")])
        if r:
            elegir(r)

    zona.bind("<Button-1>", buscar)
    if dnd:
        zona.drop_target_register(DND_FILES)
        zona.dnd_bind("<<Drop>>", lambda e: elegir(root.tk.splitlist(e.data)[0]))
    texto_zona()

    # — opciones —
    fila = ttk.Frame(marco); fila.pack(fill="x", pady=4)
    ttk.Label(fila, text="Norma:").pack(side="left")
    nombres = {NORMAS[k].nombre: k for k in NORMAS}
    cb = ttk.Combobox(fila, values=list(nombres), state="readonly", width=34)
    cb.current(0); cb.pack(side="left", padx=8)

    fila2 = ttk.Frame(marco); fila2.pack(fill="x", pady=4)
    ttk.Label(fila2, text="Encabezado corrido (opcional):").pack(side="left")
    v_enc = tk.StringVar(); ttk.Entry(fila2, textvariable=v_enc, width=30).pack(side="left", padx=8)
    fila3 = ttk.Frame(marco); fila3.pack(fill="x", pady=4)
    ttk.Label(fila3, text="Apellido (solo MLA):").pack(side="left")
    v_ape = tk.StringVar(); ttk.Entry(fila3, textvariable=v_ape, width=20).pack(side="left", padx=8)

    v_ord = tk.BooleanVar(); v_num = tk.BooleanVar(); v_tit = tk.BooleanVar(value=True); v_pdf = tk.BooleanVar()
    ttk.Checkbutton(marco, text="Ordenar referencias alfabéticamente", variable=v_ord).pack(anchor="w")
    ttk.Checkbutton(marco, text="Numerar referencias si faltan (IEEE / Vancouver)", variable=v_num).pack(anchor="w")
    ttk.Checkbutton(marco, text="Detectar títulos automáticamente si no tienen estilo", variable=v_tit).pack(anchor="w")
    ttk.Checkbutton(marco, text="Generar también un PDF (requiere LibreOffice)", variable=v_pdf).pack(anchor="w")

    boton = ttk.Button(marco, text="Reformatear documento")
    boton.pack(pady=10, ipadx=10, ipady=4)
    registro = tk.Text(marco, height=12, state="disabled", wrap="word", font=("Consolas", 9))
    registro.pack(fill="both", expand=True)

    def escribir(msg):
        registro.config(state="normal"); registro.insert("end", msg + "\n"); registro.see("end")
        registro.config(state="disabled")

    def trabajo(ruta, norma, op):
        try:
            res = reformatear(ruta, norma, opciones=op, log=lambda m: cola.put(("log", m)))
            cola.put(("ok", res))
        except ErrorReformateo as e:
            cola.put(("err", str(e)))
        except Exception as e:  # error inesperado: se muestra sin cerrar la aplicación
            cola.put(("err", f"Error inesperado ({type(e).__name__}): {e}"))

    def iniciar():
        if not estado["archivo"]:
            messagebox.showinfo("Falta el documento", "Arrastre o seleccione primero un documento.")
            return
        op = Opciones(encabezado=v_enc.get(), apellido=v_ape.get(), ordenar_refs=v_ord.get(),
                      numerar_refs=v_num.get(), detectar_titulos=v_tit.get(), exportar_pdf=v_pdf.get())
        boton.config(state="disabled")
        registro.config(state="normal"); registro.delete("1.0", "end"); registro.config(state="disabled")
        threading.Thread(target=trabajo, args=(estado["archivo"], nombres[cb.get()], op), daemon=True).start()

    def sondear():
        try:
            while True:
                tipo, dato = cola.get_nowait()
                if tipo == "log":
                    escribir("• " + dato)
                elif tipo == "ok":
                    boton.config(state="normal")
                    escribir("\n" + dato.texto())
                    if messagebox.askyesno("Listo", f"Documento creado:\n{dato.salida}\n\n¿Abrir la carpeta?"):
                        abrir_carpeta(dato.salida)
                else:
                    boton.config(state="normal")
                    escribir("✖ " + dato)
                    messagebox.showerror("No se pudo reformatear", dato)
        except queue.Empty:
            pass
        root.after(150, sondear)

    boton.config(command=iniciar)
    sondear()
    root.mainloop()


# ════════════════════════════════════════════════════════════════════════════
# 10. LÍNEA DE COMANDOS
# ════════════════════════════════════════════════════════════════════════════
def main(argv=None):
    ap = argparse.ArgumentParser(description="Reformatea un documento académico (.docx/.pdf) según una norma.")
    ap.add_argument("entrada", nargs="?", help="documento de entrada (.docx, .pdf)")
    ap.add_argument("-n", "--norma", default="apa7", choices=list(NORMAS), help="norma (por defecto apa7)")
    ap.add_argument("-o", "--salida", help="archivo .docx de salida")
    ap.add_argument("--papel", choices=["carta", "a4", "original"], help="tamaño de papel (por defecto el de la norma)")
    ap.add_argument("--encabezado", default="", help="texto de encabezado corrido")
    ap.add_argument("--apellido", default="", help="apellido para el encabezado MLA")
    ap.add_argument("--ordenar-refs", action="store_true", help="ordenar referencias alfabéticamente")
    ap.add_argument("--numerar-refs", action="store_true", help="numerar referencias (IEEE/Vancouver)")
    ap.add_argument("--sin-detectar-titulos", action="store_true", help="no detectar títulos por su aspecto")
    ap.add_argument("--columnas", type=int, choices=[1, 2], help="número de columnas (IEEE suele usar 2)")
    ap.add_argument("--pdf", action="store_true", help="generar también PDF (requiere LibreOffice)")
    ap.add_argument("--gui", action="store_true", help="abrir la interfaz gráfica")
    ap.add_argument("--listar-normas", action="store_true", help="mostrar las normas disponibles")
    a = ap.parse_args(argv)

    if a.listar_normas:
        for k, v in NORMAS.items():
            print(f"{k:10s} {v.nombre:32s} {v.fuente} {v.tam:g} pt · interlineado {v.interlineado:g} · "
                  f"márgenes (sup/inf/izq/der) {v.margenes} cm")
        return 0
    if a.gui or not a.entrada:
        lanzar_gui()
        return 0

    op = Opciones(papel=a.papel, encabezado=a.encabezado, apellido=a.apellido, ordenar_refs=a.ordenar_refs,
                  numerar_refs=a.numerar_refs, detectar_titulos=not a.sin_detectar_titulos,
                  columnas=a.columnas, exportar_pdf=a.pdf)
    try:
        res = reformatear(a.entrada, a.norma, a.salida, op, log=lambda m: print("•", m))
    except ErrorReformateo as e:
        print(f"\nError: {e}", file=sys.stderr)
        return 1
    except Exception as e:
        print(f"\nError inesperado ({type(e).__name__}): {e}", file=sys.stderr)
        return 2
    print("\n" + res.texto())
    return 0


if __name__ == "__main__":
    sys.exit(main())
